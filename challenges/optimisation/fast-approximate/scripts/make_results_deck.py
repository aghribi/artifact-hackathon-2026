# Copy of the forecasting kickoff deck with the team's results inserted after each tier's slide.
# New slides reuse the deck's own shapes (accent bar, tier tag, title, subtitle, bullet text, footer, page number),
# cloned from the tier slides, so they keep its fonts and colours; figures are PNGs from results/ or decoded from
# the notebooks. Page numbers of every slide are rewritten after the insertion.
# Usage: python -I make_results_deck.py <Forecasting.pptx> <out.pptx> <slides spec .py>
#   the spec file defines SLIDES = [{'tier': 'easy'|'medium'|'hard', 'title', 'subtitle', 'bullets': [...],
#   'bullet_box': (x, y, w, h) inches, or 'bullet_boxes': [((x, y, w, h), [...]), ...], 'bullet_size': pt, 'figs': [(png, (x, y, w, h)), ...], 'captions': [(text, (x, y, w, h)), ...],
#   'table': {'rows': [[...], ...], 'box': (x, y, w, h), 'size': pt, 'first': first column fraction}}, ...]
import copy, sys
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Inches, Pt

TIER_SLIDE = {'easy': 18, 'medium': 19, 'hard': 20}      # 0-based index of each tier's slide in the original deck
TIER_TAG = {'easy': 'TIER 1 · EASY · RESULTS', 'medium': 'TIER 2 · MEDIUM · RESULTS', 'hard': 'TIER 3 · HARD · RESULTS'}
A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'


def shape(slide, name):
    return next(s for s in slide.shapes if s.name == name)


def page_numbers(slide):
    return [sh for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip().isdigit()
            and abs(Emu(sh.top or 0).inches - 6.83) < 0.1]


def set_text(sp, lines, size=None):
    # replace the paragraphs of a cloned text shape, keeping the first paragraph's formatting for every line
    txb = sp.find('.//' + '{http://schemas.openxmlformats.org/presentationml/2006/main}txBody')
    paras = txb.findall(A + 'p')
    proto = paras[0]
    for p in paras:
        txb.remove(p)
    for line in lines:
        p = copy.deepcopy(proto)
        runs = p.findall(A + 'r')
        for r in runs[1:]:
            p.remove(r)
        runs[0].find(A + 't').text = line
        if size:
            for rp in p.iter(A + 'rPr'):
                rp.set('sz', str(int(size * 100)))
            for rp in p.iter(A + 'endParaRPr'):
                rp.set('sz', str(int(size * 100)))
            ln = p.find(f'{A}pPr/{A}lnSpc/{A}spcPts')
            if ln is not None:
                ln.set('val', str(int(size * 140)))
        txb.append(p)


def place(sp, box):
    off, ext = sp.find(f'.//{A}off'), sp.find(f'.//{A}ext')
    x, y, w, h = (Inches(v) for v in box)
    off.set('x', str(x)); off.set('y', str(y)); ext.set('cx', str(w)); ext.set('cy', str(h))


def fit_picture(slide, png, box):
    # largest size with the image's aspect ratio inside box, centred
    x, y, w, h = box
    iw, ih = Image.open(png).size
    s = min(w / iw, h / ih)
    slide.shapes.add_picture(png, Inches(x + (w - iw * s) / 2), Inches(y + (h - ih * s) / 2), Inches(iw * s), Inches(ih * s))


def table(slide, spec):
    rows = spec['rows']
    x, y, w, h = spec['box']
    t = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(h)).table
    first = spec.get('first', 0.3)                                       # row labels get a wide first column
    t.columns[0].width = Inches(w * first)
    for j in range(1, len(rows[0])):
        t.columns[j].width = Inches(w * (1 - first) / (len(rows[0]) - 1))
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = t.cell(i, j)
            c.text = str(v)
            c.fill.solid()
            c.fill.fore_color.rgb = RGBColor.from_string('E6E0D2' if i == 0 else 'FFFDF8')
            c.margin_left = c.margin_right = Inches(0.05)
            c.margin_top = c.margin_bottom = Inches(0.02)
            for p in c.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(spec.get('size', 11))
                    r.font.name = 'IBM Plex Sans'
                    r.font.bold = i == 0 or j == 0
                    r.font.color.rgb = RGBColor.from_string('14213D')


def main():
    src, out, spec_file = sys.argv[1:4]
    spec = {'__file__': spec_file}
    exec(compile(open(spec_file).read(), spec_file, 'exec'), spec)      # the spec is our own file, not downloaded data
    prs = Presentation(src)
    orig = list(prs.slides)
    sub_proto = shape(orig[1], 'Text 3')._element                        # subtitle style ('The dose and ...')
    ids = prs.slides._sldIdLst
    offset = next(int(sh.text_frame.text) - i - 1 for i, s in enumerate(orig) for sh in page_numbers(s))
    inserted = {k: 0 for k in TIER_SLIDE}
    for sd in spec['SLIDES']:
        tier = orig[TIER_SLIDE[sd['tier']]]
        s = prs.slides.add_slide(prs.slide_layouts[1])
        for ph in list(s.placeholders):
            ph._element.getparent().remove(ph._element)
        tree = s.shapes._spTree
        names = [sh.name for sh in tier.shapes]
        title_name = next(sh.name for sh in tier.shapes if sh.is_placeholder)
        tag_box, tag_text = names[1], names[2]                            # rounded tag and its text
        footer, page = names[-2], names[-1]
        bullets_name = next(sh.name for sh in tier.shapes if sh.has_text_frame and '•' in sh._element.xml)
        clone = {n: copy.deepcopy(shape(tier, n)._element) for n in (names[0], tag_box, tag_text, title_name, footer, page, bullets_name)}
        for n in (names[0], tag_box, tag_text, title_name, footer, page):
            tree.append(clone[n])
        set_text(clone[tag_text], [TIER_TAG[sd['tier']]])
        place(clone[tag_box], (0.89, 0.89, 3.1, 0.31)); place(clone[tag_text], (1.0, 0.93, 3.0, 0.26))
        set_text(clone[title_name], [sd['title']])
        if sd.get('subtitle'):
            sub = copy.deepcopy(sub_proto)
            tree.append(sub)
            set_text(sub, [sd['subtitle']])
            place(sub, (0.89, 1.95, 11.9, 0.35))
        boxes = sd.get('bullet_boxes', [(sd.get('bullet_box', (0.95, 2.5, 5.6, 4.0)), sd['bullets'])] if sd.get('bullets') else [])
        for box, lines in boxes:
            b = copy.deepcopy(clone[bullets_name])
            tree.append(b)
            set_text(b, lines, sd.get('bullet_size'))
            place(b, box)
        for text, box in sd.get('captions', []):
            c = copy.deepcopy(sub_proto)
            tree.append(c)
            set_text(c, [text], 10)
            place(c, box)
        for png, box in sd.get('figs', []):
            fit_picture(s, png, box)
        if sd.get('table'):
            table(s, sd['table'])
        # move after the tier slide and the slides already inserted there
        el = ids[-1]
        ids.remove(el)
        ids.insert(TIER_SLIDE[sd['tier']] + 1 + sum(inserted[t] for t in TIER_SLIDE if TIER_SLIDE[t] <= TIER_SLIDE[sd['tier']]), el)
        inserted[sd['tier']] += 1
    for i, s in enumerate(prs.slides):                                    # page numbers, with the deck's own offset
        for sh in page_numbers(s):
            set_text(sh._element, [f'{i + 1 + offset:02d}'])
    prs.save(out)
    print(f'{len(spec["SLIDES"])} slides inserted -> {out}')


if __name__ == '__main__':
    main()
